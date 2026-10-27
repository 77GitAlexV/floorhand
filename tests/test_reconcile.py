"""Re-uploading the manager's sheet: the sheet wins, through a visible difference,
and never makes a second copy of an order."""
import io
import random
from datetime import date

import pytest

from floorhand import db, domain, reconcile, seed
from floorhand.domain import DomainError
from floorhand.features import orders as O

TODAY = date(2026, 10, 27)
HEAD = "Order No;Customer;Product;Fabric;Frame;Qty;Price;Due;Shipped"


@pytest.fixture
def conn(tmp_path, monkeypatch):
    monkeypatch.setattr(domain, "today", lambda conn: TODAY.isoformat())
    c = db.connect(tmp_path / "r.db")
    seed.build(c, random.Random(7), today=TODAY)
    return c


def m(conn):
    return conn.execute("SELECT id FROM people WHERE role = 'manager'").fetchone()[0]


def sheet(*rows, head=HEAD):
    return "\n".join([head, *rows])


ROWS = ["901;Urban Loft;Cozy armchair;velvet beige;walnut;2;380;2026-11-20;*",
        "902;Casa Verde;Round pouf;boucle milk;;3;120;25.11;*",
        "903;Studio Blanc;Lord armchair;linen sand;oak;1;520;2026-12-01;*"]


def upload(conn, text, book="Book A", decisions=None, apply=True):
    p = reconcile.propose(conn, m(conn), text, "x.csv", book)
    if not apply:
        return p, None
    return p, reconcile.apply(conn, m(conn), p["id"], decisions or {})


def by(p, kind):
    return [r for r in p["rows"] if r["kind"] == kind]


def order_of(conn, book, ref):
    row = conn.execute("SELECT order_id FROM order_extra WHERE ext_book = ? AND ext_ref = ?",
                       (book, ref)).fetchone()
    return row and conn.execute("SELECT * FROM orders WHERE id = ?", (row[0],)).fetchone()


def count_orders(conn):
    return conn.execute("SELECT COUNT(*) FROM orders").fetchone()[0]


def test_same_sheet_twice_makes_no_duplicates(conn):
    before = count_orders(conn)
    p, res = upload(conn, sheet(*ROWS))
    assert len(by(p, "new")) == 3 and res["applied"] == 3
    assert count_orders(conn) == before + 3
    p2, _ = upload(conn, sheet(*ROWS), apply=False)
    assert "Nothing changed" in p2["message"] and len(by(p2, "same")) == 3
    assert count_orders(conn) == before + 3
    o = order_of(conn, "Book A", "902")
    assert (o["qty"], o["due_date"], o["variant"]) == (3, "2026-11-25", "boucle milk")


def test_the_same_numbers_in_another_book_are_other_orders(conn):
    upload(conn, sheet(*ROWS), book="Chairs")
    p, _ = upload(conn, sheet(*ROWS), book="Sofas", apply=False)
    assert len(by(p, "new")) == 3


def test_changes_are_shown_and_safe_ones_preselected(conn):
    upload(conn, sheet(*ROWS))
    rows = [ROWS[0].replace(";2;380;", ";3;380;"), ROWS[1].replace("25.11", "28.11"), ROWS[2]]
    p, res = upload(conn, sheet(*rows))
    changed = by(p, "changed")
    assert {c["field"] for r in changed for c in r["changes"]} == {"qty", "due_date"}
    assert all(r["default"] == "apply" for r in changed)
    assert order_of(conn, "Book A", "901")["qty"] == 3
    assert order_of(conn, "Book A", "902")["due_date"] == "2026-11-28"
    hist = conn.execute("SELECT source, reason FROM order_changes WHERE field = 'qty' "
                        "ORDER BY id DESC").fetchone()
    assert hist["source"] == "sheet" and "Book A" in hist["reason"]


def test_an_empty_cell_erases_nothing(conn):
    upload(conn, sheet(*ROWS))
    p, _ = upload(conn, sheet("901;Urban Loft;Cozy armchair;velvet beige;walnut;;380;;*",
                              *ROWS[1:]))
    assert "Nothing changed" in p["message"]
    o = order_of(conn, "Book A", "901")
    assert (o["qty"], o["due_date"]) == (2, "2026-11-20")


def test_one_number_on_two_different_rows_is_skipped_whole(conn):
    upload(conn, sheet(*ROWS))
    p, _ = upload(conn, sheet(*ROWS, ROWS[0].replace(";2;380;", ";5;380;")))
    (c,) = by(p, "conflict")
    assert c["ref"] == "901" and "different content" in c["words"]
    assert not by(p, "missing")  # the conflicting order is not reported missing
    assert order_of(conn, "Book A", "901")["qty"] == 2


def test_identical_repeated_rows_count_once(conn):
    p, _ = upload(conn, sheet(*ROWS, ROWS[1]))
    assert len(by(p, "new")) == 3 and not by(p, "conflict")


def test_renumbered_row_is_asked_never_applied_silently(conn):
    upload(conn, sheet(*ROWS))
    shifted = "903;Hygge House;Footstool;boucle ash;;4;90;2026-12-03;*"
    p, res = upload(conn, sheet(*ROWS[:2], shifted))
    (s,) = by(p, "suspect")
    assert s["default"] is None and "renumbered" in s["words"]
    assert order_of(conn, "Book A", "903")["product_id"] == \
        conn.execute("SELECT id FROM products WHERE name = 'Lord armchair'").fetchone()[0]
    p, res = upload(conn, sheet(*ROWS[:2], shifted), decisions=None, apply=False)
    res = reconcile.apply(conn, m(conn), p["id"], {by(p, "suspect")[0]["key"]: "new"})
    new = order_of(conn, "Book A", "903")
    assert new["qty"] == 4 and "Created" in res["details"][0]
    old = conn.execute("SELECT * FROM orders WHERE id = ?", (s["order_id"],)).fetchone()
    assert old["cancelled"] == 0 and old["id"] != new["id"]  # the old order lives on, unlinked


def test_missing_rows_are_listed_and_kept(conn):
    upload(conn, sheet(*ROWS))
    p, res = upload(conn, sheet(*ROWS[:2]))
    (miss,) = by(p, "missing")
    assert miss["default"] == "keep"
    assert order_of(conn, "Book A", "903")["cancelled"] == 0
    p, _ = upload(conn, sheet(*ROWS[:2]), apply=False)
    reconcile.apply(conn, m(conn), p["id"], {by(p, "missing")[0]["key"]: "cancel"})
    assert order_of(conn, "Book A", "903")["cancelled"] == 1


def test_cancel_word_in_the_sheet_cancels_on_yes(conn):
    upload(conn, sheet(*ROWS))
    o = order_of(conn, "Book A", "902")
    p, res = upload(conn, sheet(ROWS[0], ROWS[1].replace(";*", ";Cancelled, to stock"), ROWS[2]))
    (c,) = by(p, "cancel")
    assert c["default"] is None and order_of(conn, "Book A", "902")["cancelled"] == 0
    p, _ = upload(conn, sheet(ROWS[0], ROWS[1].replace(";*", ";Cancelled, to stock"), ROWS[2]),
                  apply=False)
    reconcile.apply(conn, m(conn), p["id"], {by(p, "cancel")[0]["key"]: "cancel"})
    assert domain.order_status(conn, o["id"]) == "cancelled"


def test_an_order_typed_by_hand_before_the_sheet_is_linked_not_duplicated(conn):
    cust = conn.execute("SELECT id FROM customers WHERE name = 'Urban Loft'").fetchone()[0]
    prod = conn.execute("SELECT id FROM products WHERE name = 'Cozy armchair'").fetchone()[0]
    oid = domain.create_order(conn, product_id=prod, qty=2, customer_id=cust,
                              due_date="2026-11-20", variant="velvet beige · walnut",
                              source="manual")
    before = count_orders(conn)
    p, res = upload(conn, sheet(ROWS[0]))
    (link,) = by(p, "link")
    assert link["order_id"] == oid and link["default"] == "link"
    assert count_orders(conn) == before
    assert order_of(conn, "Book A", "901")["id"] == oid


def test_sheet_without_numbers_matches_by_content(conn):
    head = "Customer,Product,Fabric,Qty,Due"
    rows = ["Nest Interiors,Daybed Lina,chenille grey,1,2026-11-30",
            "Olive Market,Square pouf,velvet emerald,2,2026-11-15"]
    before = count_orders(conn)
    upload(conn, sheet(*rows, head=head))
    p, _ = upload(conn, sheet(rows[0].replace("2026-11-30", "2026-12-04"), rows[1], head=head))
    assert count_orders(conn) == before + 2
    (ch,) = by(p, "changed")
    assert ch["changes"][0]["field"] == "due_date"


def test_stale_sheet_does_not_revert_a_hand_edit(conn):
    upload(conn, sheet(*ROWS))
    o = order_of(conn, "Book A", "901")
    O.edit(conn, m(conn), o["id"], "due_date", "2026-11-24", reason="customer called")
    p, res = upload(conn, sheet(*ROWS))  # the old sheet again
    (ch,) = by(p, "changed")
    assert ch["stale"] and ch["default"] == "skip"
    assert order_of(conn, "Book A", "901")["due_date"] == "2026-11-24"


def test_change_touching_done_work_needs_an_explicit_yes(conn):
    upload(conn, sheet(*ROWS))
    o = order_of(conn, "Book A", "901")
    u1, u2 = [r[0] for r in conn.execute("SELECT id FROM units WHERE order_id = ?", (o["id"],))]
    for uid in (u1, u2):
        cut = conn.execute("SELECT id FROM people WHERE stage_id = (SELECT id FROM stages "
                           "WHERE code = 'cut') LIMIT 1").fetchone()[0]
        domain.mark_done(conn, cut, uid)
    text = sheet(ROWS[0].replace(";2;380;", ";1;380;"), *ROWS[1:])
    p, _ = upload(conn, text)
    (ch,) = by(p, "changed")
    assert ch["danger"] and ch["default"] is None
    assert order_of(conn, "Book A", "901")["qty"] == 2
    p, _ = upload(conn, text, apply=False)
    reconcile.apply(conn, m(conn), p["id"], {by(p, "changed")[0]["key"]: "apply"})
    o = order_of(conn, "Book A", "901")
    assert o["qty"] == 1
    moved = [u for u in (u1, u2) if domain._unit_order(conn, u)["id"] != o["id"]]
    assert len(moved) == 1  # the cut item went to stock with its work
    assert domain._unit_order(conn, moved[0])["customer_id"] is None
    assert domain.done_stage_ids(conn, moved[0])


def test_unknown_names_are_never_guessed(conn):
    upload(conn, sheet(*ROWS))
    p, _ = upload(conn, sheet(ROWS[0].replace("Cozy armchair", "Cozy chair XL"), *ROWS[1:],
                              "904;Urban Loft;Throne;velvet beige;;1;;;*"), apply=False)
    (blocked,) = by(p, "blocked")
    assert "not in the catalog" in blocked["words"]
    (check,) = by(p, "check")
    assert check["choices"] == ["skip"]


def test_decisions_must_be_offered_and_apply_closes(conn):
    p, _ = upload(conn, sheet(*ROWS), apply=False)
    key = by(p, "new")[0]["key"]
    with pytest.raises(DomainError, match="pick one"):
        reconcile.apply(conn, m(conn), p["id"], {key: "cancel"})
    reconcile.apply(conn, m(conn), p["id"], {})
    with pytest.raises(DomainError, match="already applied"):
        reconcile.apply(conn, m(conn), p["id"], {})


def test_one_failing_row_does_not_stop_the_rest(conn, monkeypatch):
    real = reconcile._create

    def flaky(conn_, O_, pid, book, r, reason):
        if r["ref"] == "902":
            raise DomainError("disk full")
        return real(conn_, O_, pid, book, r, reason)
    monkeypatch.setattr(reconcile, "_create", flaky)
    p, res = upload(conn, sheet(*ROWS))
    assert res["applied"] == 2 and "disk full" in res["message"]
    assert order_of(conn, "Book A", "903") is not None


def test_xlsx_months_with_shifted_columns_template_totals_and_tails(conn):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Эталон"
    wb.active.append(["Дата", "№ заказа", "Покупатель", "Изделие", "Ткань", "Каркас", "Кол-во"])
    a = wb.create_sheet("Октябрь")
    a.append(["Дата", "№ заказа", "Покупатель", "Изделие", "Ткань", "Магазин", "Каркас",
              "Кол-во", "Цена", "Срок изготовления", "Дата отправки", "Оплачено"])
    a.append([date(2026, 10, 20), 1200, "Urban Loft", "Cozy armchair", "velvet beige", "Supplier",
              "walnut", 2, 380, date(2026, 11, 20), "*", 0])
    a.append([None, "1201 / 33", "Casa Verde", "Round pouf", "boucle milk", "", "", 3, 120,
              "05.01 urgent", "8/10", 0])
    a.append([None, None, "Итого", None, None, None, None, None, 500])
    b = wb.create_sheet("Ноябрь")
    b.append(["Дата", "№ заказа", "Покупатель", "Изделие", "Ткань", "Магазин", "Каркас",
              "Кол-во", "Цена", "Срок изготовления", "Готовность", "Дата отправки"])
    b.append([date(2026, 11, 2), 1202, "Studio Blanc", "Lord armchair", "linen sand", "", "oak",
              1, 520, date(2026, 12, 1), "готов", "*"])
    b.append(["Дата", "№ заказа", "Покупатель", "Изделие", "Ткань"])
    b.append([None, 1203, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    p = reconcile.propose(conn, m(conn), buf.getvalue(), "orders.xlsx")
    new = by(p, "new")
    assert [r["ref"] for r in new] == ["1200", "1201", "1202"]
    pouf = new[1]
    assert pouf["sheet"]["due_date"] == "2027-01-05"  # no year: the next 5 January
    assert "ref tail: 33" in pouf["note"] and "urgent" in pouf["note"]
    assert new[2]["sheet"]["qty"] == 1 and new[2]["sheet"]["variant"] == "linen sand · oak"
    whys = {s["why"] for s in p["skipped"]}
    assert {"template sheet", "total row", "header repeated"} <= whys
    assert any("1203" in w for w in whys)
    assert p["book"] == "orders"


def test_demo_sheet_shows_every_kind_of_difference(conn):
    data = reconcile.demo_sheet(conn)
    p = reconcile.propose(conn, m(conn), data, "Orders 2026.xlsx")
    kinds = {r["kind"] for r in p["rows"]}
    assert {"new", "changed", "cancel", "suspect", "missing", "conflict", "same"} <= kinds
    reconcile.apply(conn, m(conn), p["id"], {})
    again = reconcile.propose(conn, m(conn), data, "Orders 2026.xlsx")
    assert not by(again, "new")
