"""The manager's sheet, uploaded again and again, without duplicate orders.

The sheet is the source of truth, but the system never applies it blindly.
Lessons from a real workshop that ran its orders from a spreadsheet for a year:

- Rows are matched to orders by the manager's own order number (scoped to the
  sheet "book": two sheets may reuse the same numbers), or by content when the
  sheet has no number column. An order typed in by hand before the sheet came
  is linked, not duplicated.
- One number on two rows with different content is a conflict: that number is
  skipped as a whole and named in the result. Identical repeated rows count once.
- When a known number suddenly has a different product AND customer, rows were
  probably renumbered after an insert. That is asked, never applied silently
  (silent application made "hybrid" orders: product of one row, customer of another).
- An empty cell erases nothing. Columns are found by their header, on every
  sheet separately (a column inserted in later months shifted everything).
- Two steps: first the difference (new / changed / cancelled / missing), then
  "apply" on Yes. Safe changes are pre-selected; changes that touch work
  already done must be accepted explicitly; impossible ones say why.
- Three-way check: the last applied sheet values are remembered per order. If
  the sheet still says what it said last time but the order was changed in the
  system since, the sheet is stale on that field: the system value is kept
  unless the manager picks the sheet.
- A row missing from the sheet is listed, never closed automatically (a "close
  by due date" rule once killed live orders).
- Every row gets an outcome; one bad row never stops the rest.
- "Shipped" in the sheet is reported, never applied: items leave only through
  the shipping screen.
"""
import csv
import difflib
import io
import json
import re
import sqlite3
from datetime import date, datetime, timedelta

from . import domain, intake
from .domain import DomainError

MAX_ROWS = 5000
TEMPLATE_SHEETS = ("template", "sample", "эталон", "шаблон")

# header word lists; checked in this order so "Ship date" is not "Date"
IGNORE = ("price", "sum", "total", "amount", "paid", "debt", "payment", "delivery", "cost",
          "цена", "сумма", "оплат", "долг", "доставка", "погрузка", "магазин", "готовность")
HEADERS = [
    ("shipped", ("ship date", "shipped", "sent", "отправ", "отгруж")),
    ("status", ("status", "статус", "cancel", "отмен")),
    ("due", ("due", "deadline", "ship by", "deliver by", "срок", "отдать")),
    ("ref", ("order no", "order #", "order number", "ref", "№", "номер")),
    ("date", ("date", "дата")),
    ("product", ("product", "item", "model", "article", "изделие", "модель", "товар")),
    ("qty", ("qty", "quantity", "pcs", "count", "кол", "шт")),
    ("customer", ("customer", "client", "buyer", "store", "покупатель", "клиент", "заказчик")),
    ("fabric", ("fabric", "colour", "color", "variant", "upholstery", "ткань", "цвет")),
    ("frame", ("frame", "legs", "wood", "каркас", "ножки", "комплек")),
    ("note", ("note", "comment", "remark", "примеч", "коммент")),
]
EXACT_REF = ("#", "no", "no.", "n", "№")
CANCEL_WORDS = ("cancel", "отмен", "отказ", "refus")
STOCK_WORDS = ("stock", "свободн", "склад", "free")
TOTAL_WORDS = ("total", "итого", "всего")


# --- reading the file -------------------------------------------------------

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).replace(" ", " ").strip()


def _header_key(h: str) -> str | None:
    h = re.sub(r"\s+", " ", h.strip().lower().replace("-", ""))
    if not h:
        return None
    if h in EXACT_REF:
        return "ref"
    if any(w in h for w in IGNORE):
        return None
    for key, words in HEADERS:
        if any(w in h for w in words):
            return key
    return None


def _header_map(cells: list[str]) -> dict[str, int] | None:
    cols: dict[str, int] = {}
    for i, c in enumerate(cells):
        k = _header_key(c)
        if k and k not in cols:
            cols[k] = i
    if len(cols) >= 2 and ("product" in cols or "ref" in cols):
        return cols
    return None


def read_sheets(data: bytes | str, filename: str = "") -> list[tuple[str, list[list[str]]]]:
    """[(sheet name, rows of cell text)]. .xlsx (every sheet) or CSV/TSV text."""
    if isinstance(data, bytes) and data[:2] == b"PK":
        try:
            import openpyxl
        except ImportError:  # pragma: no cover
            raise DomainError("Reading .xlsx needs openpyxl on the server.")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise DomainError("This .xlsx file could not be opened. Save it again and retry.")
        out = []
        for ws in wb.worksheets:
            rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
            out.append((ws.title, rows))
        return out
    text = data.decode("utf-8-sig", errors="replace") if isinstance(data, bytes) else data
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        raise DomainError("The file is empty.")
    delim = max("\t;,|", key=lines[0].count)
    return [("table", [[_cell(c) for c in r]
                       for r in csv.reader(io.StringIO("\n".join(lines)), delimiter=delim)])]


# --- understanding one row --------------------------------------------------

def _parse_ref(text: str) -> tuple[str, str]:
    """(main ref, tail). '1234 / 56', '1234-Smith', '1234 (56)', '1234*' -> '1234'."""
    t = text.strip()
    if not t:
        return "", ""
    m = re.match(r"^(\d{1,7})(?:\.0)?\s*(.*)$", t)
    if m:
        tail = m[2].strip(" */-()№#").strip()
        return m[1], tail
    return re.sub(r"\s+", " ", t.lower()), ""


def parse_date(text: str, today: date) -> tuple[str | None, str]:
    """(ISO date or None, the rest of the text). A date with no year is the
    nearest one ahead (a sheet row from late December says '05.01')."""
    t = text.strip()
    if not t:
        return None, ""
    if m := re.search(r"\b(20\d\d)-(\d\d)-(\d\d)\b", t):
        try:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat(), \
                (t[:m.start()] + t[m.end():]).strip(" ,;")
        except ValueError:
            return None, t
    if m := re.search(r"\b(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?\b", t):
        d, mo, y = int(m[1]), int(m[2]), m[3]
        try:
            if y:
                year = int(y) + (2000 if len(y) == 2 else 0)
                found = date(year, mo, d)
            else:
                found = date(today.year, mo, d)
                if found < today - timedelta(days=60):
                    found = found.replace(year=today.year + 1)
            return found.isoformat(), (t[:m.start()] + t[m.end():]).strip(" ,;")
        except ValueError:
            return None, t
    d = intake._parse_date(t, today)
    return (d.isoformat(), "") if d else (None, t)


def _qty(text: str) -> int | None:
    m = re.search(r"\d{1,4}", text or "")
    return int(m[0]) if m else None


def _ship_state(text: str) -> dict:
    """What the sheet says about shipping/cancelling. Read, never applied:
    shipping happens only on the shipping screen."""
    low = text.lower().strip()
    if not low or low == "*":
        return {"state": "open"}
    if any(w in low for w in CANCEL_WORDS):
        return {"state": "cancel_to_stock" if any(w in low for w in STOCK_WORDS) else "cancel",
                "text": text}
    if m := re.fullmatch(r"\s*(\d{1,3})\s*/\s*(\d{1,3})\s*", low):
        return {"state": "partial", "sent": int(m[1]), "of": int(m[2]), "text": text}
    return {"state": "shipped", "text": text}


def _rows(sheets, today: date) -> tuple[list[dict], list[dict]]:
    """Data rows with their sheet address, and the rows skipped with why."""
    rows, skipped = [], []
    for name, cells in sheets:
        if name.strip().lower() in TEMPLATE_SHEETS:
            skipped.append({"where": name, "why": "template sheet"})
            continue
        cols = None
        for i, r in enumerate(cells, 1):
            where = f"{name}!{i}"
            if not any(r):
                continue
            hm = _header_map(r)
            if hm:
                if cols is not None:
                    skipped.append({"where": where, "why": "header repeated"})
                cols = hm
                continue
            if cols is None:
                continue
            get = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else ""  # noqa: E731
            if not get("product") and any(c.lower().startswith(TOTAL_WORDS) for c in r if c):
                skipped.append({"where": where, "why": "total row"})
                continue
            ref, tail = _parse_ref(get("ref"))
            if not get("product") and not get("customer"):
                if ref:
                    skipped.append({"where": where, "why": f"number {ref} reserved, row empty"})
                continue
            due, due_rest = parse_date(get("due"), today)
            status = _ship_state(get("status")) if get("status") else _ship_state(get("shipped"))
            if status["state"] == "open" and get("status"):
                status = _ship_state(get("shipped"))
            note = " · ".join(x for x in (tail and f"ref tail: {tail}", due_rest, get("note")) if x)
            rows.append({"where": where, "ref": ref, "product_text": get("product"),
                         "qty_text": get("qty"), "customer_text": get("customer"),
                         "fabric_text": get("fabric"), "frame_text": get("frame"),
                         "due": due, "due_text": get("due"), "ship": status, "note": note,
                         "has": {k: bool(get(k)) for k in ("product", "qty", "customer",
                                                           "fabric", "frame", "due")}})
    return rows, skipped


def _resolve(row: dict, cat: dict) -> dict:
    """Catalog ids for one row. Unknown names stay as text with suggestions."""
    out = {}
    if row["has"]["product"]:
        pid, opts = intake._match("product", row["product_text"], cat)
        out["product"] = {"id": pid, "text": row["product_text"],
                          "value": cat["product"].get(pid), "options": opts}
    if row["has"]["customer"]:
        t = row["customer_text"]
        if t.strip().lower() in ("stock", "to stock", "на склад", "склад"):
            out["customer"] = {"id": None, "stock": True, "text": t, "value": "to stock"}
        else:
            cid, opts = intake._match("customer", t, cat)
            out["customer"] = {"id": cid, "stock": False, "text": t,
                               "value": cat["customer"].get(cid), "options": opts}
    if row["has"]["qty"]:
        out["qty"] = _qty(row["qty_text"])
    if row["has"]["fabric"] or row["has"]["frame"]:
        text = f"{row['fabric_text']} {row['frame_text']}"
        v = intake._variant(text, cat)
        known = (v["fabric"] or not row["has"]["fabric"]) and (v["frame"] or not row["has"]["frame"])
        out["variant"] = {"value": " · ".join(x for x in (v["fabric"], v["frame"]) if x),
                          "known": bool(known), "text": text.strip(),
                          "options": v["fabric_options"]}
    if row["due"]:
        out["due_date"] = row["due"]
    return out


# --- the difference -----------------------------------------------------------

def _order_values(conn, oid: int) -> dict:
    o = conn.execute("SELECT o.*, p.name AS product, c.name AS customer FROM orders o "
                     "JOIN products p ON p.id = o.product_id "
                     "LEFT JOIN customers c ON c.id = o.customer_id WHERE o.id = ?",
                     (oid,)).fetchone()
    return {"product": o["product"], "product_id": o["product_id"], "qty": o["qty"],
            "customer": o["customer"] or "to stock", "customer_id": o["customer_id"],
            "variant": o["variant"] or "", "due_date": o["due_date"], "number": o["number"],
            "cancelled": bool(o["cancelled"]), "id": o["id"]}


def _sheet_values(res: dict) -> dict:
    """What the row says, in the same shape as _order_values (known values only)."""
    out = {}
    if res.get("product", {}).get("id"):
        out["product"] = res["product"]["value"]
    c = res.get("customer")
    if c and (c["id"] or c.get("stock")):
        out["customer"] = c["value"]
    if res.get("qty"):
        out["qty"] = res["qty"]
    if res.get("variant", {}).get("known"):
        out["variant"] = res["variant"]["value"]
    if res.get("due_date"):
        out["due_date"] = res["due_date"]
    return out


def _consequence(conn, oid: int, field: str, new) -> tuple[str, str]:
    """('safe'|'danger'|'blocked', words) for changing one field from the sheet."""
    from .features import orders as O
    us = O.units_state(conn, oid)
    shipped = [u for u in us if u["shipped"]]
    began = [u for u in us if u["started"] and not u["shipped"]]
    if field == "qty":
        cur = len(us)
        if new >= cur:
            return "safe", f"+{new - cur} item(s)" if new > cur else ""
        if new < len(shipped):
            return "blocked", f"{len(shipped)} already shipped"
        idle = len([u for u in us if not u["started"] and not u["shipped"]])
        drop = cur - new
        if drop <= idle:
            return "safe", f"{drop} not-started item(s) removed"
        return "danger", f"{drop - idle} item(s) in work go to stock without customer"
    if field in ("product", "variant", "customer") and shipped:
        return "blocked", f"{len(shipped)} pc already shipped"
    if field == "product" and began:
        return "danger", (f"{len(began)} started item(s) stay as they are and go to stock; "
                          "new ones start at the first stage")
    if field == "variant" and began:
        return "danger", (f"{len(began)} item(s) already in work in the old variant; "
                          "send them back from the order card if needed")
    if field == "customer" and began:
        return "safe", "printed sheets show the old customer: print again"
    return "safe", ""


FIELDS = ("product", "qty", "customer", "variant", "due_date")


def _changes(conn, oid: int, sheet: dict, last: dict, res: dict) -> list[dict]:
    have = _order_values(conn, oid)
    out = []
    for f in FIELDS:
        if f not in sheet:
            continue
        new, old = sheet[f], have[f]
        if str(new) == str(old or ""):
            continue
        level, words = _consequence(conn, oid, f, new)
        stale = f in last and str(last[f]) == str(new)  # the sheet did not change this field
        out.append({"field": f, "old": old, "new": new, "level": level, "words": words,
                    "stale": stale})
    for f, key in (("product", "product"), ("customer", "customer")):
        r = res.get(key)
        if r and not r.get("id") and not r.get("stock"):
            hint = ", ".join(o["label"] for o in r.get("options", [])) or "none close"
            out.append({"field": f, "old": have[f], "new": r["text"], "level": "blocked",
                        "words": f"“{r['text']}” is not in the catalog (closest: {hint})",
                        "stale": False})
    v = res.get("variant")
    if v and not v["known"]:
        out.append({"field": "variant", "old": have["variant"], "new": v["text"],
                    "level": "blocked",
                    "words": "fabric not in the catalog"
                             + (f" (closest: {', '.join(v['options'])})" if v["options"] else ""),
                    "stale": False})
    return out


def _content_key(sheet: dict) -> tuple:
    return (sheet.get("customer"), sheet.get("product"), sheet.get("variant", ""))


def propose(conn: sqlite3.Connection, person_id: int, data: bytes | str, filename: str = "",
            book: str = "") -> dict:
    """Read the sheet and work out the difference. Writes nothing but the
    proposal (an agent_actions row holding the whole difference and the rows
    as read, a snapshot of this upload)."""
    who = domain.person(conn, person_id)
    if who["role"] not in ("manager", "owner"):
        raise DomainError("Only the office can upload the orders sheet.")
    book = (book or re.sub(r"\.(xlsx|csv|tsv|txt)$", "", filename or "", flags=re.I)
            or "orders sheet").strip()[:80]
    today = date.fromisoformat(domain.today(conn))
    rows, skipped = _rows(read_sheets(data, filename), today)
    if not rows:
        raise DomainError("No order rows found. The sheet needs a header row with at least "
                          "“Product” and one more column (Qty, Customer, Order No …).")
    if len(rows) > MAX_ROWS:
        raise DomainError(f"The sheet has {len(rows)} rows; the limit is {MAX_ROWS}. Split it.")
    cat = intake._catalog(conn)
    for r in rows:
        r["res"] = _resolve(r, cat)
        r["sheet"] = _sheet_values(r["res"])

    linked = {}
    for e in conn.execute("SELECT x.order_id, x.ext_ref, x.fingerprint FROM order_extra x "
                          "JOIN orders o ON o.id = x.order_id WHERE x.ext_book = ?", (book,)):
        linked[e["ext_ref"] or f"@{e['order_id']}"] = {
            "order_id": e["order_id"], "last": json.loads(e["fingerprint"] or "{}")}
    by_order = {v["order_id"]: k for k, v in linked.items()}

    # group rows by number: identical repeats count once, different content conflicts
    out_rows, seen_orders = [], set()
    groups: dict[str, list[dict]] = {}
    for r in rows:
        if r["ref"]:
            groups.setdefault(r["ref"], []).append(r)
    singles = []
    for ref, rs in groups.items():
        contents = {json.dumps([r["sheet"], r["ship"]["state"]], sort_keys=True) for r in rs}
        if len(contents) > 1:
            out_rows.append({"key": f"ref:{ref}", "ref": ref, "where": ", ".join(r["where"] for r in rs),
                             "kind": "conflict", "choices": ["skip"], "default": "skip",
                             "words": f"Number {ref} is on {len(rs)} rows with different content "
                                      f"({', '.join(r['where'] for r in rs)}). Skipped as a whole: "
                                      "fix the sheet so each number is one order."})
            if ref in linked:
                seen_orders.add(linked[ref]["order_id"])
            continue
        r = rs[0]
        if len(rs) > 1:
            r["note"] = (r["note"] + " · " if r["note"] else "") + f"same row {len(rs)} times"
        singles.append(r)
    singles += [r for r in rows if not r["ref"]]

    # orders the office typed before the sheet: candidates to link by content
    unlinked = [dict(o) for o in conn.execute(
        "SELECT o.id FROM orders o LEFT JOIN order_extra x ON x.order_id = o.id "
        "WHERE o.cancelled = 0 AND (x.ext_book IS NULL) ORDER BY o.number")]
    unlinked_vals = {o["id"]: _order_values(conn, o["id"]) for o in unlinked}
    no_ref_linked = [k for k in linked if k.startswith("@")]

    for r in singles:
        key = f"ref:{r['ref']}" if r["ref"] else f"row:{r['where']}"
        sheet = r["sheet"]
        base = {"key": key, "ref": r["ref"], "where": r["where"], "sheet": sheet,
                "note": r["note"], "ship": r["ship"], "row": {
                    k: r[k] for k in ("product_text", "qty_text", "customer_text",
                                      "fabric_text", "frame_text", "due_text")}}
        target = None
        if r["ref"] and r["ref"] in linked:
            target = linked[r["ref"]]
        elif not r["ref"]:  # no number column: match by content among this book's orders
            for k in no_ref_linked:
                oid = linked[k]["order_id"]
                if oid in seen_orders:
                    continue
                if _content_key(_order_values(conn, oid)) == _content_key(sheet) or \
                        _content_key(linked[k]["last"]) == _content_key(sheet):
                    target = linked[k]
                    break
        if target:
            oid = target["order_id"]
            seen_orders.add(oid)
            have = _order_values(conn, oid)
            have["status"] = domain.order_status(conn, oid)
            row = {**base, "order_id": oid, "number": have["number"]}
            if r["ship"]["state"] in ("cancel", "cancel_to_stock"):
                if have["cancelled"]:
                    out_rows.append({**row, "kind": "same", "choices": [], "default": None,
                                     "words": "Cancelled in the sheet and in the system."})
                elif have["status"] == "shipped":
                    out_rows.append({**row, "kind": "blocked", "choices": ["skip"],
                                     "default": "skip",
                                     "words": f"The sheet says “{r['ship']['text']}”, but "
                                              f"#{have['number']} is already shipped. Check "
                                              "the sheet: a return is a new order."})
                else:
                    out_rows.append({**row, "kind": "cancel", "choices": ["cancel", "skip"],
                                     "default": None,
                                     "words": f"The sheet says “{r['ship']['text']}”. Cancel "
                                              f"#{have['number']}: finished items go to stock "
                                              "without customer, the rest stops."})
                continue
            prod_changed = "product" in sheet and sheet["product"] != have["product"]
            cust_changed = "customer" in sheet and sheet["customer"] != have["customer"]
            last = target["last"]
            if prod_changed and cust_changed and not (
                    last.get("product") == sheet.get("product")
                    and last.get("customer") == sheet.get("customer")):
                out_rows.append({**row, "kind": "suspect", "choices": ["apply", "new", "skip"],
                                 "default": None, "changes": _changes(conn, oid, sheet, last,
                                                                      r["res"]),
                                 "words": f"Number {r['ref']} was {have['product']} for "
                                          f"{have['customer']} (#{have['number']}); now the row "
                                          f"says {sheet['product']} for {sheet['customer']}. "
                                          "Were rows renumbered? Update the order, make a new "
                                          "one, or skip."})
                continue
            changes = _changes(conn, oid, sheet, last, r["res"])
            out_rows.append(_changed_row(row, changes, r, have))
            continue
        # new to this book: maybe the office already typed it in by hand
        match = None
        if "customer" in sheet and "product" in sheet:
            for o in unlinked:
                v = unlinked_vals[o["id"]]
                if o["id"] in seen_orders or v["customer"] != sheet["customer"] or \
                        v["product"] != sheet["product"]:
                    continue
                if sheet.get("variant", v["variant"]) == v["variant"] and (
                        sheet.get("due_date") == v["due_date"] or sheet.get("qty") == v["qty"]):
                    match = v
                    break
        if match:
            seen_orders.add(match["id"])
            changes = _changes(conn, match["id"], sheet, {}, r["res"])
            out_rows.append({**base, "kind": "link", "order_id": match["id"],
                             "number": match["number"], "changes": changes,
                             "choices": ["link", "new", "skip"], "default": "link",
                             "words": f"Looks like #{match['number']} typed in earlier "
                                      f"({match['product']} for {match['customer']}). Link it "
                                      "to this row instead of making a second order"
                                      + (", and take the sheet's values." if changes else ".")})
            continue
        out_rows.append(_new_row(base, r, cat))

    for k, v in linked.items():
        oid = v["order_id"]
        if oid in seen_orders:
            continue
        have = _order_values(conn, oid)
        status = domain.order_status(conn, oid)
        if status in ("shipped", "cancelled"):
            continue
        out_rows.append({"key": f"missing:{oid}", "ref": None if k.startswith("@") else k,
                         "where": "", "kind": "missing", "order_id": oid,
                         "number": have["number"], "choices": ["keep", "cancel"],
                         "default": "keep",
                         "words": f"#{have['number']} ({have['product']} for {have['customer']}, "
                                  f"{status.replace('_', ' ')}) is not in this upload. Kept as "
                                  "it is unless you cancel it."})

    summary = _summary(out_rows, skipped)
    payload = {"book": book, "filename": filename, "rows": out_rows, "skipped": skipped,
               "summary": summary}
    cur = conn.execute(
        "INSERT INTO agent_actions (kind, payload, model, created_at, for_person) "
        "VALUES ('sheet', ?, 'rules', ?, ?)",
        (json.dumps(payload), domain.now(conn).isoformat(timespec="seconds"), person_id))
    conn.commit()
    return {"id": cur.lastrowid, **payload, "message": summary["words"]}


def _changed_row(row: dict, changes: list[dict], r: dict, have: dict) -> dict:
    ship = r["ship"]
    info = []
    if have.get("status") == "shipped":
        pass
    elif ship["state"] == "shipped":
        info.append(f"The sheet says shipped (“{ship['text']}”); ship it on the Shipping screen "
                    "if it really left.")
    elif ship["state"] == "partial":
        info.append(f"The sheet says {ship['sent']} of {ship['of']} shipped.")
    if not changes:
        return {**row, "kind": "same", "choices": [], "default": None, "changes": [],
                "info": info, "words": "Same as the system."}
    live = [c for c in changes if c["level"] != "blocked"]
    if not live:
        return {**row, "kind": "blocked", "changes": changes, "choices": ["skip"],
                "default": "skip", "info": info,
                "words": "; ".join(f"{c['field']}: {c['words']}" for c in changes) + "."}
    stale = all(c["stale"] for c in live)
    danger = any(c["level"] == "danger" for c in live)
    words = []
    if stale:
        words.append("Changed in the system since the last upload; the sheet still has the "
                     "old value. Kept as in the system unless you take the sheet.")
    if danger:
        words.append("Touches work already done: accept it explicitly.")
    return {**row, "kind": "changed", "changes": changes,
            "choices": ["apply", "skip"],
            "default": "skip" if stale else (None if danger else "apply"),
            "danger": danger, "stale": stale, "info": info, "words": " ".join(words)}


def _new_row(base: dict, r: dict, cat: dict) -> dict:
    res, sheet = r["res"], r["sheet"]
    problems = []
    p = res.get("product")
    if not p or not p.get("id"):
        hint = ", ".join(o["label"] for o in (p or {}).get("options", []))
        problems.append(f"product “{(p or {}).get('text', '')}” is not in the catalog"
                        + (f" (closest: {hint})" if hint else ""))
    v = res.get("variant")
    if v and not v["known"]:
        problems.append(f"fabric “{v['text']}” is not in the catalog")
    if r["ship"]["state"] in ("cancel", "cancel_to_stock"):
        return {**base, "kind": "skip", "choices": [], "default": None,
                "words": "Cancelled in the sheet and never made here: nothing to do."}
    if r["ship"]["state"] == "shipped":
        problems.append("the sheet says it is already shipped, but it was never in the system")
    if problems:
        return {**base, "kind": "check", "choices": ["skip"], "default": "skip",
                "words": "Cannot create: " + "; ".join(problems) + ". Fix the sheet (or the "
                         "catalog) and upload again."}
    c = res.get("customer")
    new_customer = bool(c and not c["id"] and not c.get("stock") and c["text"])
    qty = sheet.get("qty")
    words = f"New order: {qty or 1} × {sheet['product']}"
    words += f" for {c['text'] if new_customer else sheet.get('customer', 'stock')}"
    if new_customer:
        words += " (a new customer)"
    if sheet.get("due_date"):
        words += f", due {sheet['due_date']}"
    if not qty:
        words += ". Quantity is empty: 1 is used"
    return {**base, "kind": "new", "choices": ["create", "skip"], "default": "create",
            "new_customer": c["text"] if new_customer else None, "words": words + "."}


def _summary(rows: list[dict], skipped: list[dict]) -> dict:
    count = {}
    for r in rows:
        count[r["kind"]] = count.get(r["kind"], 0) + 1
    actionable = [r for r in rows if r.get("choices")]
    if not actionable or all(r["kind"] in ("same",) for r in rows):
        words = f"Nothing changed: {count.get('same', 0)} row(s) match the system."
    else:
        parts = [f"{count[k]} {label}" for k, label in (
            ("new", "new"), ("changed", "changed"), ("link", "to link"), ("cancel", "cancelled"),
            ("suspect", "renumbered?"), ("missing", "missing"), ("conflict", "conflicting"),
            ("check", "to check"), ("blocked", "blocked"), ("same", "same")) if count.get(k)]
        words = "Sheet read: " + ", ".join(parts) + "."
    return {"count": count, "skipped": len(skipped), "words": words}


# --- apply ------------------------------------------------------------------

def apply(conn: sqlite3.Connection, person_id: int, action_id: int,
          decisions: dict[str, str]) -> dict:
    """Apply the difference the person accepted. Each decision must be one of
    the row's offered choices; rows without a decision use their default (a
    dangerous or suspicious row has none, so it is left alone). One failing
    row is reported and the rest go on."""
    from .features import orders as O
    row = conn.execute("SELECT * FROM agent_actions WHERE id = ? AND kind = 'sheet'",
                       (action_id,)).fetchone()
    if row is None:
        raise DomainError("No such sheet upload.")
    if row["status"] != "proposed":
        raise DomainError("This upload is already applied. Upload the sheet again to compare.")
    payload = json.loads(row["payload"])
    book = payload["book"]
    keys = {r["key"] for r in payload["rows"]}
    for k, v in decisions.items():
        if k not in keys:
            raise DomainError(f"Unknown row {k}.")
        r = next(x for x in payload["rows"] if x["key"] == k)
        if v not in r.get("choices", []):
            raise DomainError(f"Row {r['where'] or k}: pick one of {', '.join(r['choices'])}.")
    reason = f"sheet “{book}”"
    done, errors, untouched = [], [], 0
    for r in payload["rows"]:
        choice = decisions.get(r["key"], r.get("default"))
        if not choice or choice in ("skip", "keep"):
            untouched += bool(r.get("choices"))
            r["outcome"] = "left as is"
            continue
        try:
            msg = _apply_row(conn, O, person_id, book, r, choice, reason)
            r["outcome"] = msg
            done.append(msg)
        except DomainError as exc:
            conn.rollback()
            r["outcome"] = f"not applied: {exc}"
            errors.append(f"{r['where'] or r['key']}: {exc}")
    conn.execute("UPDATE agent_actions SET payload = ?, status = 'applied', decided_by = ?, "
                 "decided_at = ?, result = ? WHERE id = ?",
                 (json.dumps(payload), person_id, domain.now(conn).isoformat(timespec="seconds"),
                  json.dumps({"done": len(done), "errors": errors}), action_id))
    conn.commit()
    msg = f"Applied {len(done)} row(s)." if done else "Nothing applied."
    if untouched:
        msg += f" {untouched} left as they are."
    if errors:
        msg += f" {len(errors)} could not be applied: " + "; ".join(errors[:5])
    return {"ok": True, "applied": len(done), "errors": errors, "details": done, "message": msg}


def _remember(conn, O, oid: int, book: str, ref: str | None, sheet: dict) -> None:
    """What the sheet said at this upload: the base of the next three-way check."""
    O.set_extra(conn, oid, ext_book=book, ext_ref=ref or None,
                fingerprint=json.dumps(sheet, sort_keys=True))


def _apply_row(conn, O, person_id: int, book: str, r: dict, choice: str, reason: str) -> str:
    kind = r["kind"]
    if kind == "missing" and choice == "cancel":
        O.cancel(conn, person_id, r["order_id"], f"not in {reason} any more")
        return f"#{r['number']} cancelled (missing from the sheet)."
    if kind == "cancel":
        res = O.cancel(conn, person_id, r["order_id"], f"{reason}: {r['ship']['text']}")
        _remember(conn, O, r["order_id"], book, r["ref"], r["sheet"])
        conn.commit()
        return res["message"]
    if (kind == "new" and choice == "create") or (kind in ("suspect", "link") and choice == "new"):
        if kind == "suspect":  # the old order keeps its content, loses the number
            O.set_extra(conn, r["order_id"], ext_ref=None)
            O.log(conn, r["order_id"], person_id, "sheet_ref", r["ref"], None,
                  f"{reason}: number now used by another row", "sheet")
        return _create(conn, O, person_id, book, r, reason)
    if kind in ("changed", "suspect", "link") and choice in ("apply", "link"):
        oid = r["order_id"]
        msgs = []
        for c in r.get("changes", []):
            if c["level"] == "blocked":
                msgs.append(f"{c['field']} not changed: {c['words']}")
                continue
            res = _edit_from_sheet(conn, O, person_id, oid, c, r, reason)
            if res.get("ok") is False:
                raise DomainError(res["message"])
            msgs.append(res["message"])
        _remember(conn, O, oid, book, r["ref"], r["sheet"])
        if kind == "link":
            O.log(conn, oid, person_id, "sheet_ref", None, f"{book} {r['ref'] or r['where']}",
                  "linked to the sheet row", "sheet")
        conn.commit()
        return f"#{r['number']}: " + (" ".join(msgs) if msgs else "linked to the sheet.")
    raise DomainError(f"“{choice}” does not fit this row.")


def _edit_from_sheet(conn, O, person_id, oid, c, r, reason) -> dict:
    f, new = c["field"], c["new"]
    if f == "product":
        return O.edit(conn, person_id, oid, "product", new, reason=reason, started="to_stock",
                      source="sheet")
    if f == "qty":
        return O.edit(conn, person_id, oid, "qty", new, reason=reason, started="to_stock",
                      source="sheet")
    if f == "variant":
        return O.edit(conn, person_id, oid, "variant", new, reason=reason, keep_work=True,
                      source="sheet")
    if f == "customer":
        cust = "stock" if new == "to stock" else new
        return O.edit(conn, person_id, oid, "customer", cust, reason=reason, source="sheet")
    return O.edit(conn, person_id, oid, "due_date", new, reason=reason, source="sheet")


def _create(conn, O, person_id: int, book: str, r: dict, reason: str) -> str:
    sheet = r["sheet"]
    cat = intake._catalog(conn)
    pid, _ = intake._match("product", sheet["product"], cat)
    cust = None
    if r.get("new_customer"):
        cust = conn.execute("INSERT INTO customers (name) VALUES (?)",
                            (r["new_customer"],)).lastrowid
    elif sheet.get("customer") not in (None, "to stock"):
        cust, _ = intake._match("customer", sheet["customer"], cat)
    oid = domain.create_order(conn, product_id=pid, qty=sheet.get("qty") or 1,
                              customer_id=cust, due_date=sheet.get("due_date"),
                              variant=sheet.get("variant", ""), source="sheet", commit=False)
    O.set_extra(conn, oid, note=r.get("note") or "", ext_book=book, ext_ref=r["ref"] or None,
                fingerprint=json.dumps(sheet, sort_keys=True))
    O.log(conn, oid, person_id, "created", None, f"{book} {r['ref'] or r['where']}", reason,
          "sheet")
    conn.commit()
    number = conn.execute("SELECT number FROM orders WHERE id = ?", (oid,)).fetchone()[0]
    return f"Created #{number} from row {r['where']}."


# --- demo: the manager's sheet ----------------------------------------------

DEMO_BOOK = "Orders 2026"


def link_demo_orders(conn: sqlite3.Connection, n: int = 40) -> None:
    """Demo data: the latest open customer orders came from the manager's
    sheet; give them the manager's own numbers (different from ours)."""
    from .features import orders as O
    rows = conn.execute("SELECT id FROM orders WHERE customer_id IS NOT NULL AND cancelled = 0 "
                        "ORDER BY number DESC LIMIT ?", (n,)).fetchall()
    for i, r in enumerate(reversed(rows)):
        v = _order_values(conn, r["id"])
        sheet = {k: v[k] for k in FIELDS if v[k] not in (None, "")}
        O.set_extra(conn, r["id"], ext_book=DEMO_BOOK, ext_ref=str(2401 + i),
                    fingerprint=json.dumps(sheet, sort_keys=True))


def demo_sheet(conn: sqlite3.Connection) -> bytes:
    """The manager's sheet as it looks a day later: a quantity and a due date
    changed, a fabric changed on an order in work, a cancellation, a missing
    row, two new rows (one for a new customer), one number on two rows, a
    renumbered row, a repeated header, a total row and a template sheet -
    everything the real sheet does. All names are the demo's own."""
    import openpyxl
    from .features import orders as O
    rows = conn.execute("SELECT x.order_id, x.ext_ref FROM order_extra x JOIN orders o ON "
                        "o.id = x.order_id WHERE x.ext_book = ? AND o.cancelled = 0 "
                        "ORDER BY CAST(x.ext_ref AS INTEGER)", (DEMO_BOOK,)).fetchall()
    head = ["Date", "Order No", "Customer", "Product", "Fabric", "Frame", "Qty", "Price",
            "Due", "Shipped", "Paid", "Note"]
    wb = openpyxl.Workbook()
    tpl = wb.active
    tpl.title = "Template"
    tpl.append(head)
    oct_ws, nov_ws = wb.create_sheet("October"), wb.create_sheet("November")
    oct_ws.append(head)
    nov_ws.append(["Date", "Order No", "Customer", "Product", "Fabric", "Frame", "Qty",
                   "Ready", "Price", "Due", "Shipped", "Paid", "Note"])  # a column inserted
    today = date.fromisoformat(domain.today(conn))
    in_work = next((r["order_id"] for r in rows if any(
        u["started"] and not u["shipped"] for u in O.units_state(conn, r["order_id"]))), None)
    vals, open_idx = [], []
    for i, r in enumerate(rows):
        v = _order_values(conn, r["order_id"])
        fabric, _, frame = (v["variant"] or "").partition(" · ")
        shipped = domain.order_status(conn, r["order_id"]) == "shipped"
        rec = {"ref": r["ext_ref"], "customer": v["customer"], "product": v["product"],
               "fabric": fabric, "frame": frame, "qty": v["qty"], "due": v["due_date"],
               "shipped": "*" if not shipped else today.strftime("%d.%m.%Y"), "note": ""}
        if not shipped and r["order_id"] != in_work:
            open_idx.append(i)
        vals.append((r["order_id"], rec))
    if len(open_idx) < 8:
        raise DomainError("Not enough open demo orders linked to the sheet.")
    pick = [vals[i][1] for i in open_idx]
    # the day's edits, the way people make them
    pick[0]["qty"] += 1
    if pick[1]["due"]:
        pick[1]["due"] = (date.fromisoformat(pick[1]["due"]) + timedelta(days=4)
                          ).strftime("%d.%m")
    for oid, rec in vals:
        if oid == in_work:
            rec["fabric"] = "boucle ash" if rec["fabric"] != "boucle ash" else "linen sand"
            break
    pick[2]["shipped"] = "Cancelled — to stock"
    twin = dict(pick[3])
    twin["qty"] = twin["qty"] + 2
    other = next(r for r in pick[5:] if r["product"] != pick[4]["product"]
                 and r["customer"] != pick[4]["customer"])
    pick[4].update(customer=other["customer"], product=other["product"],
                   fabric=other["fabric"], frame=other["frame"])
    pick[6]["ref"] = pick[6]["ref"] + " / 17"  # the store's own number in the tail
    missing = vals.pop(open_idx[7])
    half = len(vals) // 2
    for ws, part in ((oct_ws, vals[:half]), (nov_ws, vals[half:])):
        for oid, rec in part:
            row = [today.strftime("%d.%m.%Y"), rec["ref"], rec["customer"], rec["product"],
                   rec["fabric"], rec["frame"], rec["qty"], 0, rec["due"] or "",
                   rec["shipped"], 0, rec["note"]]
            if ws is nov_ws:
                row.insert(7, "")
            ws.append(row)
    oct_ws.append(["", twin["ref"], twin["customer"], twin["product"], twin["fabric"],
                   twin["frame"], twin["qty"], 0, twin["due"] or "", "*", 0, ""])
    oct_ws.append(["", "", "Total", "", "", "", "", 0, "", "", 0, ""])
    nov_ws.append(head[:7] + ["Ready"] + head[7:])  # the header repeated mid-sheet
    top = max(int(re.match(r"\d+", r["ext_ref"])[0]) for r in rows)
    due = (today + timedelta(days=12)).strftime("%d.%m")
    nov_ws.append([today.strftime("%d.%m.%Y"), str(top + 1), "Hygge House", "Round pouf",
                   "boucle milk", "", 3, "", 0, due, "*", 0, "with handles"])
    nov_ws.append([today.strftime("%d.%m.%Y"), str(top + 2), "Bluebird Cafe", "Bruno chair",
                   "linen sand", "black", 6, "", 0, due, "*", 0, ""])
    nov_ws.append(["", str(top + 3), "", "", "", "", "", "", "", "", "", "", ""])  # reserved
    del missing  # left out of the sheet on purpose
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
